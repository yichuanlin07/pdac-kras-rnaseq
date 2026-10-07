import sys, copy, re, json, zipfile, hashlib
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

base, template, target = map(Path, sys.argv[1:4])
target.parent.mkdir(parents=True, exist_ok=True)
res = base / "results"
figs = base / "figures"
read = lambda n: pd.read_csv(res / n, sep="\t")
sc = read("scenario_summary.tsv")
ps = read("paper_comparison_summary.tsv")
sen = read("sensitivity_full.tsv")
pairs = read("candidate_pair_directions.tsv")
validation = json.loads((res / "independent_numerical_validation.json").read_text())
doc = Document(template)
patterns = [copy.deepcopy(p._p) for p in doc.paragraphs]
body = doc._element.body
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)


def para(text="", pattern=4, cite=None):
    el = copy.deepcopy(patterns[pattern])
    rp = None
    runs = el.findall(qn("w:r"))
    if runs:
        rp = copy.deepcopy(runs[0].find(qn("w:rPr")))
    for child in list(el):
        if child.tag != qn("w:pPr"):
            el.remove(child)
    body.insert(len(body) - 1, el)
    p = Paragraph(el, doc._body)
    r = p.add_run(text)
    if rp is not None:
        r._r.insert(0, rp)
    spacing = p.paragraph_format
    if pattern == 4:
        spacing.space_after = Pt(16)
    elif pattern == 18:
        spacing.space_after = Pt(18)
    elif pattern in (3, 15):
        spacing.space_before = Pt(16)
        spacing.space_after = Pt(8)
    elif pattern == 17:
        spacing.space_before = Pt(4)
        spacing.space_after = Pt(8)
    elif pattern == 16:
        spacing.space_before = Pt(4)
        spacing.space_after = Pt(8)
    elif pattern == 34:
        spacing.space_after = Pt(8)
        contextual = p._p.get_or_add_pPr().find(qn("w:contextualSpacing"))
        if contextual is None:
            contextual = OxmlElement("w:contextualSpacing")
            p._p.get_or_add_pPr().append(contextual)
        contextual.set(qn("w:val"), "0")
    if cite:
        r = p.add_run(str(cite))
        r.font.name = "Times New Roman"
        r.font.size = Pt(8)
        r.font.superscript = True
    return p


def heading(text, page=False):
    return para(text, 15 if page else 3)


def cap(label, text, cite=None):
    para(label, 17)
    para(text, 18, cite)


def picture(name, width=6.72):
    p = para("", 16)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(figs / (name + ".png")), width=Inches(width))


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    if widths:
        for i, w in enumerate(widths):
            t.columns[i].width = Inches(w)
    for i, s in enumerate(headers):
        t.rows[0].cells[i].text = s
    for row in rows:
        cells = t.add_row().cells
        for i, s in enumerate(row):
            cells[i].text = str(s)
    for row_number, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement("w:cantSplit"))
        for i, c in enumerate(row.cells):
            if widths:
                c.width = Inches(widths[i])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "FFFFFF")
            tcpr.append(shade)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "bottom", "left", "right"):
                margin = OxmlElement("w:" + edge)
                margin.set(qn("w:w"), "120")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            tcpr.append(margins)
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.keep_with_next = row_number < len(t.rows) - 1
                if re.fullmatch("[+-]?[\\d,]+(?:[.]\\d+)?(?:e[+-]?\\d+)?", c.text):
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0, 0, 0)
    for c in t.rows[0].cells:
        for r in c.paragraphs[0].runs:
            r.bold = True
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "bottom", "insideH", "insideV", "left", "right"]:
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "000000")
        borders.append(e)
    t._tbl.tblPr.append(borders)
    para("", 18).paragraph_format.space_after = Pt(8)
    return t


def sci(x):
    return f"{x:.2g}"


num = lambda x: f"{int(x):,}"
para("KRAS suppression in pancreatic cancer cells", 0)
para("WEEK4 PROGRESS REPORT", 1)
para("Yichuan Lin", 2)
heading("1 Overview")
para(
    "I compared DESeq2 and edgeR using the real RNA-seq counts from eight pancreatic cancer cell lines treated with KRAS or control siRNA. Both methods recovered a shared expression response, with differences in the genes passing the significance threshold. The published study provides the biological context.",
    cite=1,
)
heading("2 Work completed")
para(
    "The 28,395-gene matrix matched the 17 original run tables and sample metadata. Counts were nonnegative integers, sample order matched, and assigned-count totals agreed. ENA identifies local P5 (SRR24828471) as MP2_K2v2 and local P5_v2 (SRR24828472) as MP2_K2; the local v2 naming is reversed.",
    cite=2,
)
para(
    "For the exploratory main analysis, I summed the counts from SRR24828471 and SRR24828472 into one MiaPaca-2 KRAS observation, following the meeting tutorial. This leaves 16 observations, one control and one treatment per cell line.",
    cite=3,
)
para(
    "Both methods use ~condition + cell_line, with Control as the reference, and retain genes with at least 10 counts in two samples. The full-rank design has nine columns and seven residual degrees of freedom. DESeq2 uses size factors and a Wald test; edgeR uses trimmed mean of M-values (TMM) normalization and robust quasi-likelihood fitting.",
    cite="4,5",
)
para(
    "At false discovery rate (FDR) < 0.05, DESeq2 detected 5,943 genes and edgeR detected 6,431, with 5,442 shared genes (Jaccard 0.785). Unshrunk log2 fold-change (log2FC) estimates correlated at Pearson r = 0.884 across 17,819 genes. Principal component analysis (PCA) showed larger differences between cell lines than within matched pairs.",
    cite=6,
)
para(
    "I also ran both methods retaining only SRR24828472, then only SRR24828471. Across these two choices and the merged analysis, 5,084 genes remained significant with the same direction in both methods. The comparisons assess sensitivity to P5 handling. The original analysis was run twice, with identical values in 81 result files; reversing the contrast changed the log2FC signs and preserved p values.",
    cite=6,
)
heading("3 Next steps")
para(
    "For the October 10 discussion, I will review the method differences, PCA and genes near the FDR threshold, and prepare the next progress update. The P5 processing follows the tutorial; the remaining question is whether the two libraries came from the same biological material or separate experiments. Leave-one-cell-line-out analysis is planned for the next session.",
    cite=3,
)
heading("4 Summary")
para(
    "The main expression response persists across methods and P5 choices, although individual gene lists change. KRAS, DUSP6, SPRY4, MYC and CCND1 decrease consistently in the fitted models; EMP2 increases. These genes provide a starting point for discussing the KRAS response and its variation across cell lines."
)
heading("Appendix A Sample structure and transformed expression", True)
picture("A1_PCA")
cap(
    "Figure A1  PCA of the main 16-sample analysis",
    "PCA uses the 500 most variable genes after DESeq2 variance-stabilizing transformation (VST; blind = FALSE) or edgeR TMM log counts per million (logCPM; prior count 2), centred without scaling. Cell-line effects remain in the data. Blue circles are Control and red triangles are KRAS siRNA; lines connect matched conditions. PC1/PC2 explain 35.1%/21.7% for DESeq2 and 33.7%/20.6% for edgeR. These components describe sample structure; treatment effects are tested in the count models.",
)
picture("A2_sample_distances")
cap(
    "Figure A2  Sample distances",
    "Euclidean distances use all 17,819 retained genes, with complete-linkage ordering. C = Control; K = KRAS siRNA; MP2 = MiaPaca-2. Darker blue indicates greater distance. Median within-line and between-line distances are 32.3 and 118.3 for VST, and 81.7 and 304.2 for logCPM. Each transformation has its own distance scale. Most matched conditions cluster closely, supporting adjustment for cell line.",
)
heading("Appendix B Differential expression and method comparison", True)
picture("B1_volcanoes")
cap(
    "Figure B1  Differential expression after KRAS siRNA",
    "Blue and red mark genes with FDR < 0.05 and negative or positive log2FC, respectively; grey marks FDR >= 0.05. The horizontal line is FDR 0.05, with no effect-size cutoff for the primary lists. Positive log2FC means higher expression after KRAS siRNA. DESeq2 displays apeglm-shrunken effect sizes with the original Wald adjusted p values; edgeR displays its reported logFC and FDR. The vertical scales differ because the tests differ.",
    7,
)
picture("B2_counts_and_agreement")
cap(
    "Figure B2  Gene counts and unshrunk effect-size agreement",
    "DESeq2 detects 2,912 downregulated and 3,031 upregulated genes; edgeR detects 2,923 and 3,508. The intersection is 5,442 and union 6,932, giving Jaccard 0.785 (intersection divided by union). The scatter uses unshrunk estimates; shared DE genes are lapis blue. Pearson r = 0.884 and Spearman rho = 0.910 across all 17,819 shared finite estimates. Axes are cropped to +/-8; correlations use all shared estimates.",
)
para(
    "The methods differ in normalization, dispersion moderation and testing. DESeq2 independent filtering at alpha = 0.05 leaves 2,764 retained genes without adjusted p values, so its FDR calculation covers 15,055 genes. edgeR adjusts across all 17,819 retained genes. Applying Benjamini-Hochberg (BH) adjustment to the common finite-p set gives 5,797 DESeq2 and 6,431 edgeR genes. The primary lists use each method’s native adjustment; the common-set calculation shows the contribution of different multiple-testing sets.",
    18,
)
heading("Appendix C P5 sensitivity and limits", True)
para("Table C1  Differential expression under alternative P5 choices", 17)
table(
    ["P5 choice", "Genes retained", "DESeq2 DE", "edgeR DE", "Both methods"],
    [
        [
            {
                "merge_p5": "Sum both (main)\nSRR24828471 + SRR24828472",
                "retain_SRR24828472": "Retain SRR24828472 only",
                "retain_SRR24828471": "Retain SRR24828471 only",
            }[r.scenario],
            num(r.prefilter_pass),
            num(r.DESeq2_sig),
            num(r.edgeR_sig),
            num(r.overlap),
        ]
        for r in sc.itertuples()
    ],
    [2.1, 1.2, 1.1, 1.1, 1.2],
)
para(
    "The main analysis combines SRR24828471 (public MP2_K2v2) and SRR24828472 (public MP2_K2) by summing each gene’s counts into one MiaPaca-2 KRAS column. This follows the meeting tutorial. Its alternate branch keeps SRR24828472 alone; the additional comparison keeps SRR24828471 alone. All original 17-run inputs are preserved.",
    18,
    3,
)
para(
    "The tutorial states that the paper did not distinguish technical from biological repeats. The runs have different sample accessions and FASTQ MD5 values, and logCPM correlation is 0.993. These observations do not identify how the libraries were generated. The three P5 choices therefore assess processing sensitivity under the tutorial’s assumptions.",
    18,
    3,
)
picture("C1_P5_sensitivity")
cap(
    "Figure C1  Effect-size sensitivity to P5 handling",
    "Comparing the merged result with SRR24828472 alone and SRR24828471 alone, DESeq2 DE-list Jaccard values are 0.936 and 0.945; edgeR values are 0.932 and 0.920. Unshrunk log2FC correlations range from 0.992 to 0.996. Axes are cropped to +/-6; correlations use all finite paired estimates. Global agreement remains high, while some genes cross the FDR threshold.",
)
heading("Appendix C continued", True)
para("Table C2  Examples requiring cautious interpretation", 17)
rows = []
for sy in ["DPM3", "EGR1", "TGFBR3"]:
    r = sen[sen.SYMBOL == sy].iloc[0]
    rows.append(
        [
            sy,
            sci(r.merge_p5_DESeq2_padj),
            sci(r.merge_p5_edgeR_FDR),
            sci(r.retain_SRR24828472_edgeR_FDR),
            sci(r.retain_SRR24828471_edgeR_FDR),
        ]
    )
table(
    ["Gene", "DESeq2 FDR, sum", "edgeR FDR, sum", "edgeR FDR, 472", "edgeR FDR, 471"],
    rows,
    [0.8, 1.5, 1.5, 1.5, 1.4],
)
para(
    "DPM3 loses edgeR significance with SRR24828472 alone. EGR1 passes the threshold in the main DESeq2 result and falls above it in edgeR; TGFBR3 falls above it in both. The additive model estimates a common treatment effect across lines. It does not estimate line-specific effects or within-line biological variability, and the P5 comparisons share observations.",
    18,
)
heading("Appendix D Benchmark against the original paper", True)
picture("D1_paper_benchmark")
cap(
    "Figure D1  Local estimates versus paper edgeR estimates",
    "Data S1 contains 14,799 Ensembl rows. Excluding 224 missing/nonnumeric Entrez IDs and all 38 rows with ambiguous numeric-ID mappings leaves 14,537 unique numeric Entrez IDs. Of these, 14,200 have finite estimates in both local methods and define the correlation set. Axes are cropped to +/-6. Both the paper and local analyses compare KRAS with control siRNA.",
    1,
)
para("Table D1  DE-list agreement on the 14,200-gene common universe", 17)
rows = []
for r in ps.itertuples():
    rows.append(
        [
            r.method,
            "FDR < .05" + (" + |LFC| > .5" if r.absLFC_gt else ""),
            num(r.local_sig),
            num(r.paper_sig),
            num(r.overlap),
            f"{r.jaccard:.3f}",
        ]
    )
table(
    ["Local method", "Selection rule", "Local DE", "Paper DE", "Overlap", "Jaccard"],
    rows,
    [1.05, 2.05, 0.9, 0.9, 0.9, 0.85],
)
para(
    "With |log2FC| > 0.5, 1,497/1,497 DESeq2 shared genes and 1,509/1,510 edgeR shared genes agree in direction with the paper. At FDR < 0.05 alone, the corresponding counts are 4,499/4,500 and 4,355/4,357. Applying both cutoffs to the full paper CSV selects 1,728 rows; 1,690 remain after unique-ID mapping and restriction to the common gene set. The effect-size cutoff is a selection rule applied after testing.",
    18,
)
para(
    "The paper used STAR, Gencode v30, Salmon and tximport; the local counts came from HISAT2 and RefSeq featureCounts. The paper also restricted gene biotypes/chromosomes, required 20 reads in two samples and used edgeR v3. Here the filter is 10 counts in two samples and edgeR is v4. The comparison measures agreement between these pipelines. DESeq2’s higher correlation alone cannot rank their accuracy.",
    18,
    1,
)
para(
    "Klomp et al. link KRAS suppression to ERK-related transcription and cell-cycle regulation. Decreases in DUSP6/SPRY4 and MYC/CCND1 fit that context. The pooled estimates identify expression changes; determining which are direct KRAS responses requires further experiments and closer examination of individual cell lines.",
    18,
    1,
)
heading("Appendix E Candidate genes for discussion", True)
picture("E1_candidate_paired_expression")
cap(
    "Figure E1  Candidate expression in matched cell lines",
    "DESeq2-normalized counts are shown as log2(count + 1). Lines connect matched conditions for each cell line; blue points are Control and red points are KRAS siRNA. MYC decreases in seven of eight lines and EMP2 increases in all eight. The plotted MiaPaca-2 KRAS observation sums both P5 libraries, as in the main model.",
)
para("Table E1  Main estimates and paired directions", 17)
rows = []
for sy in ["KRAS", "DUSP6", "SPRY4", "MYC", "CCND1", "EMP2"]:
    r = sen[sen.SYMBOL == sy].iloc[0]
    p = pairs[pairs.SYMBOL == sy].iloc[0]
    rows.append(
        [
            sy,
            f"{r.merge_p5_DESeq2_log2FC:.2f}",
            f"{r.merge_p5_edgeR_log2FC:.2f}",
            sci(r.merge_p5_DESeq2_padj),
            sci(r.merge_p5_edgeR_FDR),
            f"{p.paired_lines_down}/8 down" if sy != "EMP2" else "8/8 up",
        ]
    )
table(
    ["Gene", "DESeq2 LFC", "edgeR LFC", "DESeq2 FDR", "edgeR FDR", "Paired direction"],
    rows,
    [0.85, 1.1, 1.1, 1.2, 1.2, 1.2],
)
para(
    "All six genes remain significant with the same direction in both methods under all three P5 choices. I selected them for discussion of KRAS/ERK and cell-cycle responses, adding EMP2 as a strong upward example. KRAS checks the knockdown; DUSP6/SPRY4 and MYC/CCND1 connect the estimates to the published study. EMP2 merits follow-up because of its increase, although the present analysis does not explain its mechanism.",
    18,
    1,
)
para(
    "This is a selected gene review, with no enrichment analysis. The full DE tables retain all input genes and their filter status. The DE tests assess a zero-effect null; effect size and consistency across paired cell lines need separate consideration.",
    18,
)
heading("Appendix F Transformation and dispersion diagnostics", True)
picture("F1_mean_SD", 6.0)
cap(
    "Figure F1  Mean and standard deviation after transformation",
    "Across-sample standard deviation (SD) is plotted against mean expression for log2 normalized counts, VST and TMM logCPM. Cell-line differences remain in the SD. VST and logCPM provide the inputs for sample distances and PCA; differential expression is fitted to counts.",
)
picture("F2_dispersion_summary", 6.0)
cap(
    "Figure F2  Dispersion estimates",
    "The panels show DESeq2 final negative-binomial (NB) dispersions and the fitted trend, edgeR biological coefficient of variation (square root of NB dispersion), and moderated quasi-likelihood (QL) posterior variance with its prior trend. The main scenario’s three-page dispersion_diagnostics.pdf provides the native R plots.",
    cite="4,5",
)
para(
    "Software: R 4.5.2; DESeq2 1.50.2; edgeR 4.8.2; apeglm 1.32.0. Package versions and the dependency lockfile accompany the reproducible source package.",
    18,
)
heading("References", True)
references = [
    "Klomp JA, Klomp JE, Stalnecker CA, Bryant KL, Edwards AC, Drizyte-Miller K, et al. Defining the KRAS- and ERK-dependent transcriptome in KRAS-mutant cancers. Science. 2024;384(6700):eadk0775. doi:10.1126/science.adk0775. Data S1.",
    "European Nucleotide Archive. PRJNA980201 run metadata [Internet]. [cited 2026 Oct 4]. Available from: https://www.ebi.ac.uk/ena/browser/view/PRJNA980201.",
    "Adrian. DESeq2 Tutorial KRAS; edgeR Tutorial KRAS; weekly tasks announcement. Biomedical DryLab [private course materials]. 2026 Oct 3. Accessed 2026 Oct 4.",
    "Love MI, Huber W, Anders S. Moderated estimation of fold change and dispersion for RNA-seq data with DESeq2. Genome Biol. 2014;15(12):550. doi:10.1186/s13059-014-0550-8.",
    "Chen Y, Chen L, Lun ATL, Baldoni PL, Smyth GK. edgeR v4: powerful differential analysis of sequencing data with expanded functionality and improved support for small counts and larger datasets. Nucleic Acids Res. 2025;53(2):gkaf018. doi:10.1093/nar/gkaf018.",
    "Lin Y. PDAC KRAS RNA-seq Week 4 analysis [Internet]. GitHub; 2026 Oct 4 [cited 2026 Oct 6]. Available from: https://github.com/yichuanlin07/pdac-kras-rnaseq/tree/6f34d7fe90474f61d3519574ecccfa22d4880e41/analysis/2026-10-04-week4.",
    "Zhu A, Ibrahim JG, Love MI. Heavy-tailed prior distributions for sequence count data: removing the noise and preserving large differences. Bioinformatics. 2019;35(12):2084-92. doi:10.1093/bioinformatics/bty895.",
]
for ref in references:
    para(ref, 34)
para("Review: ____________________    Date: ____________________", 37)
doc.core_properties.title = "KRAS suppression in pancreatic cancer cells"
doc.core_properties.subject = "RNA-seq Week 4 progress report"
revision = doc.core_properties._element.find(qn("cp:revision"))
if revision is not None:
    doc.core_properties._element.remove(revision)
from datetime import datetime, timezone

doc.core_properties.modified = datetime(2026, 10, 6, tzinfo=timezone.utc)
settings = doc.settings.element
for e in settings.findall(qn("w:updateFields")):
    settings.remove(e)
update = OxmlElement("w:updateFields")
update.set(qn("w:val"), "true")
settings.append(update)
temp = target.with_suffix(".build.docx")
doc.save(temp)
editable = {
    "word/document.xml",
    "word/_rels/document.xml.rels",
    "[Content_Types].xml",
    "word/settings.xml",
    "docProps/core.xml",
}
with (
    zipfile.ZipFile(template) as old,
    zipfile.ZipFile(temp) as new,
    zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as final,
):
    oldnames = set(old.namelist())
    newnames = set(new.namelist())
    for name in old.namelist():
        final.writestr(name, new.read(name) if name in editable else old.read(name))
    for name in new.namelist():
        if name not in oldnames:
            final.writestr(name, new.read(name))
with zipfile.ZipFile(template) as old, zipfile.ZipFile(target) as final:
    same = [n for n in old.namelist() if n not in editable]
    assert all((old.read(n) == final.read(n) for n in same))
    from lxml import etree

    a = etree.fromstring(old.read("word/document.xml"))
    b = etree.fromstring(final.read("word/document.xml"))
    assert etree.tostring(a.find(".//" + qn("w:sectPr"))) == etree.tostring(
        b.find(".//" + qn("w:sectPr"))
    )
    text = final.read("word/document.xml").decode()
    assert "TOC" not in text and "bookmarkStart" not in text
temp.unlink()
qa = dict(
    template_sha256=hashlib.sha256(template.read_bytes()).hexdigest(),
    preserved_opaque_parts=len(same),
    all_preserved_byte_for_byte=True,
    section_geometry_exact=True,
    body_font="Times New Roman",
    true_superscript_citations=True,
    figures=len(doc.inline_shapes),
    tables=len(doc.tables),
    body_sections=["1 Overview", "2 Work completed", "3 Next steps", "4 Summary"],
    source_results="2026-10-04-week4",
    body_paragraph_after_pt=16,
    appendix_paragraph_after_pt=18,
    table_cell_margin_pt=6,
    horizontal_guides_in=["B2_counts_and_agreement", "E1_candidate_paired_expression"],
)
(target.parent / (target.stem + "_structure_validation.json")).write_text(json.dumps(qa, indent=2))
print(target)
